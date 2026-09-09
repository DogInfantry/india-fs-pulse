"""The whole report as one Word document, assembled from the generated memos.

The point of building it this way: `insights/*.md` are written by the analysis layer,
and every figure in them is interpolated from a computed value. Restating any of those
numbers here would create a second place for them to drift, which is exactly what
CLAUDE.md rule 1 exists to prevent. So this script writes almost no prose of its own. It
renders what the pipeline already produced, in module order, and appends the provenance
ledger the site also publishes.

Module order comes from each memo's own `generator:` front-matter field, which reads
`analysis/01_...` through `analysis/10_...`. That is already the order of the argument,
so no separate list has to be maintained and there is nothing to forget to update.

Deliberately not committed. The `.docx` is a build artefact, reproducible from the same
data as the site, and a binary that would churn on every run. `deliverables/` is ignored
for the same reason `data/raw/` is.

Run with `python run.py report`, which is a separate task rather than part of `analyze`
so that a missing python-docx cannot break the existing pipeline.
"""
from __future__ import annotations

import json
import re
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
INSIGHTS = ROOT / "insights"
SITE_DATA = ROOT / "site" / "src" / "data"
LEDGER = ROOT / "data-pipeline" / "data" / "processed" / "_provenance.json"
OUT_DIR = ROOT / "deliverables"
OUT = OUT_DIR / "india-fs-pulse.docx"

TITLE = "India FS Pulse"
SUBTITLE = ("India built the world's largest real-time payments network. "
            "Under zero MDR it earns almost nothing directly. "
            "Who captures the value, and is there an investable business model?")

# The three inline markers the memos actually use. Checked against all ten before
# being written down, so the renderer covers the real input rather than a guess.
INLINE = re.compile(r"(\*\*[^*]+\*\*|(?<!\*)\*[^*]+\*(?!\*)|`[^`]+`)")


def front_matter(path: Path) -> tuple[dict, str]:
    """Split a generated memo into its YAML block and its body."""
    text = path.read_text(encoding="utf-8")
    if not text.startswith("---"):
        return {}, text
    _, block, body = text.split("---", 2)
    meta: dict = {"sources": []}
    key = None
    for line in block.strip().splitlines():
        if line.startswith("  - "):
            if key == "sources":
                meta["sources"].append(json.loads(line[4:].strip()))
            continue
        if ":" in line:
            key, _, value = line.partition(":")
            key, value = key.strip(), value.strip()
            if value:
                meta[key] = value.strip('"')
    return meta, body


def add_runs(para, text: str) -> None:
    """Render the inline subset: bold, italic and code."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            para.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = para.add_run(piece[1:-1])
            run.font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*"):
            para.add_run(piece[1:-1]).italic = True
        else:
            para.add_run(piece)


def add_table(doc, rows: list[str]) -> None:
    """A pipe table, minus its separator row."""
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for c in cells if not all(re.fullmatch(r":?-{2,}:?", x or "-") for x in c)]
    if not cells:
        return
    width = max(len(c) for c in cells)
    table = doc.add_table(rows=0, cols=width)
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(cells):
        line = table.add_row().cells
        for j in range(width):
            para = line[j].paragraphs[0]
            add_runs(para, row[j] if j < len(row) else "")
            if i == 0:
                for run in para.runs:
                    run.bold = True


def render_markdown(doc, body: str) -> None:
    lines = body.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped or stripped.startswith("<!--"):
            i += 1
            continue
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            doc.add_paragraph()
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        else:
            # A memo paragraph is hard-wrapped across several lines; rejoin it so
            # Word does the wrapping instead of preserving the source's line breaks.
            block = []
            while i < len(lines) and lines[i].strip() and not re.match(
                    r"\s*(\||-\s|#{2,3}\s)", lines[i]):
                block.append(lines[i].strip())
                i += 1
            add_runs(doc.add_paragraph(), " ".join(block))
            continue
        i += 1


def main() -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt
    except ImportError:
        raise SystemExit("python-docx is not installed. Run: pip install python-docx")

    memos = sorted(INSIGHTS.glob("*.md"))
    if not memos:
        raise SystemExit("no memos in insights/. Run: python run.py analyze")
    meta = json.loads((SITE_DATA / "pipeline_meta.json").read_text(encoding="utf-8"))
    ledger = json.loads(LEDGER.read_text(encoding="utf-8")) if LEDGER.exists() else {}

    parsed = [(path, *front_matter(path)) for path in memos]
    # Module order, taken from the generator field rather than a list kept in step by hand.
    parsed.sort(key=lambda t: t[1].get("generator", "zz"))

    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)

    title = doc.add_heading(TITLE, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(sub, f"*{SUBTITLE}*")
    stamp = doc.add_paragraph()
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(stamp, (
        f"Generated {date.today().isoformat()} from a Python pipeline over "
        f"**{meta['publishers']} public sources**, {meta['fetchers']} fetchers and "
        f"{meta['analysis_modules']} analysis modules. Every figure below is interpolated "
        "from a computed value, so this document cannot drift from the data behind it."))

    # The answer before the contents, from the same answer.json the page reads, so
    # the document and the site cannot state a different conclusion.
    ans_path = SITE_DATA / "answer.json"
    if ans_path.exists():
        ans = json.loads(ans_path.read_text(encoding="utf-8"))
        doc.add_page_break()
        doc.add_heading("The answer", level=1)
        add_runs(doc.add_paragraph(), ans["governing"])
        for i, pil in enumerate(ans["pillars"], 1):
            doc.add_heading(f"{i}. {pil['n']}  {pil['claim']}", level=2)
            add_runs(doc.add_paragraph(), pil["support"])

    doc.add_page_break()
    doc.add_heading("Contents", level=1)
    for _, m, _ in parsed:
        add_runs(doc.add_paragraph(style="List Number"), m.get("title", "Untitled"))

    for _, m, body in parsed:
        doc.add_page_break()
        doc.add_heading(m.get("title", "Untitled"), level=1)
        note = doc.add_paragraph()
        add_runs(note, f"*Generated by {m.get('generator', 'unknown')} on "
                       f"{m.get('generated', 'unknown')}.*")
        render_markdown(doc, body)
        if m.get("sources"):
            doc.add_heading("Sources for this memo", level=2)
            for src in m["sources"]:
                add_runs(doc.add_paragraph(style="List Bullet"), src)

    doc.add_page_break()
    doc.add_heading("Appendix: data sources and provenance", level=1)
    add_runs(doc.add_paragraph(), (
        "Generated from the fetchers' provenance ledger, the same record the site "
        "publishes. Code in this repository is Apache 2.0 licensed; **each dataset "
        "retains its own publisher's terms**."))
    rows = ["| Dataset | Publisher | Coverage | Rows | Accessed | Licence |",
            "|---|---|---|---|---|---|"]
    for name in sorted(ledger):
        e = ledger[name]
        rows.append(f"| {name} | {e['publisher']} | {e['coverage']} | {e['rows']:,} | "
                    f"{e['accessed']} | {e['licence']} |")
    add_table(doc, rows)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    size_kb = OUT.stat().st_size / 1024
    print(f"   wrote {OUT.relative_to(ROOT)}  ({len(parsed)} memos, "
          f"{len(ledger)} datasets, {size_kb:,.0f} kB)")


if __name__ == "__main__":
    main()
